import json
import logging
import threading
import faiss
import numpy as np
import fitz
from docx import Document
from pathlib import Path
from sentence_transformers import SentenceTransformer
from langchain.text_splitter import RecursiveCharacterTextSplitter
from app.config import (
    CHUNK_SIZE, CHUNK_OVERLAP, ALLOWED_EXTENSIONS,
    EMBEDDING_MODEL, FAISS_INDEX_DIR, DOMAIN_COLLECTIONS,
)
from app.domain_router import detect_domain

logger = logging.getLogger(__name__)


class DocumentError(ValueError):
    """A problem with an uploaded file, with a message that is safe to show
    the user (the API turns it into an HTTP 400 response)."""


# ── Document Parsing ──


def extract_tables(page: fitz.Page) -> list[tuple[fitz.Rect, str]]:
    """Find the tables on a PDF page.
    Returns (bounding box, formatted text) for each table."""
    tables = []
    for table in page.find_tables().tables:
        rows = [" | ".join(str(cell) if cell else "" for cell in row) for row in table.extract()]
        tables.append((fitz.Rect(table.bbox), "\n[TABLE]\n" + "\n".join(rows) + "\n[/TABLE]\n"))
    return tables


def page_text_with_tables(page: fitz.Page) -> str:
    """A page's text in reading order, with each table written out once.

    Plain text extraction already returns a table's cells as loose lines
    ("Net profit\\n$6.1M\\n$5.0M"), so adding a formatted copy would store every
    table twice. Instead, text blocks that sit inside a table are skipped, and
    the formatted table goes where its first cell appeared, which keeps it
    next to its heading.
    """
    tables = extract_tables(page)
    parts = []
    placed = set()

    for x0, y0, x1, y1, text, _block_no, block_type in page.get_text("blocks"):
        if block_type != 0:  # 1 = image block, which has no text
            continue
        centre = fitz.Point((x0 + x1) / 2, (y0 + y1) / 2)
        table_number = next((i for i, (box, _) in enumerate(tables) if box.contains(centre)), None)
        if table_number is None:
            parts.append(text)
        elif table_number not in placed:
            parts.append(tables[table_number][1])
            placed.add(table_number)

    # A table with no text blocks inside its box (rare) is still included
    parts.extend(table_text for i, (_, table_text) in enumerate(tables) if i not in placed)
    return "".join(parts).strip()


def parse_pdf(file_path: str) -> list[dict]:
    """Extract text and tables per page."""
    doc = fitz.open(file_path)
    try:
        # An encrypted PDF opens without error but yields no text,
        # so check explicitly to give the user the real reason.
        if doc.needs_pass:
            raise DocumentError(
                "This PDF is password-protected. Remove the password and upload it again."
            )

        pages = []
        for page_num, page in enumerate(doc, start=1):
            text = page_text_with_tables(page)
            if text:
                pages.append({"text": text, "page": page_num})
        return pages
    finally:
        doc.close()


def parse_docx(file_path: str) -> list[dict]:
    """Extract text and tables from a DOCX file."""
    doc = Document(file_path)
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            parts.append("[TABLE]\n" + "\n".join(rows) + "\n[/TABLE]")

    if not parts:
        return []
    return [{"text": "\n\n".join(parts), "page": 1}]


def parse_txt(file_path: str) -> list[dict]:
    """Read a plain text file.

    Tries UTF-8 first (with or without a byte-order mark), then falls back to
    Windows-1252, the encoding Notepad and Excel often use on Windows.
    """
    raw = Path(file_path).read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")

    # Decoding bytes by hand keeps Windows "\r\n" line endings. Normalise them,
    # or the chunker's "\n\n" paragraph separator never matches.
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []
    return [{"text": text, "page": 1}]


def parse_document(file_path: str) -> list[dict]:
    """Route to the correct parser based on file extension."""
    ext = Path(file_path).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentError(f"Unsupported file type: {ext}. Allowed: {sorted(ALLOWED_EXTENSIONS)}")

    parser = {".pdf": parse_pdf, ".docx": parse_docx, ".txt": parse_txt}
    try:
        return parser[ext](file_path)
    except DocumentError:
        raise
    except Exception as exc:
        # Uploaded files are untrusted: a corrupted or mislabelled file can make
        # the parsing libraries raise almost anything. Log the real error for
        # debugging, and give the user a message they can act on.
        logger.warning("Could not parse %s: %r", Path(file_path).name, exc)
        kind = ext[1:].upper()
        raise DocumentError(
            f"Could not read this {kind} file. It may be corrupted, or not really a {kind} file."
        ) from exc


# ── Chunking ──


splitter = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE,
    chunk_overlap=CHUNK_OVERLAP,
    separators=["\n\n", "\n", ". ", " ", ""],
)


def chunk_document(pages: list[dict], filename: str) -> list[dict]:
    """Split parsed pages into smaller chunks with metadata."""
    chunks = []
    chunk_id = 0
    for page in pages:
        splits = splitter.split_text(page["text"])
        for split_text in splits:
            chunks.append({
                "text": split_text,
                "metadata": {
                    "filename": filename,
                    "page": page["page"],
                    "chunk_id": chunk_id,
                },
            })
            chunk_id += 1
    return chunks


# ── Embedding Model ──
# Loaded once at startup (~130MB download on first run).
# All other modules import this same instance.

embedding_model = SentenceTransformer(EMBEDDING_MODEL)
EMBEDDING_DIM = embedding_model.get_sentence_embedding_dimension()

# Read at most 256 tokens of each text (the default is 512). Every chunk and
# question in this app fits in that, and it makes embedding about 45% faster
# on a CPU, which matters when a 500-page offer document has ~2,000 chunks.
embedding_model.max_seq_length = 256


# ── FAISS Vector Store ──
# FAISS stores vectors in an efficient index for fast similarity search.
# Unlike ChromaDB, FAISS doesn't store metadata — so we keep a parallel
# JSON file (documents.json) per domain with the text and metadata for
# each vector. The position in the JSON list matches the vector's index
# in FAISS.
#
# On disk:
#   faiss_data/
#     legal_docs/
#       index.faiss      ← the vector index
#       documents.json   ← [{text, metadata}, ...] matching each vector
#     finance_docs/
#       ...

# Uploads and deletes load an index, change it, and save it back. If two ran
# at once, one could overwrite the other's changes, and a question asked
# mid-save could read half-written files. Every index read and write takes
# this lock, so they happen one at a time. It's re-entrant because
# embed_and_store() holds it while calling load_index(), which takes it too.
_index_lock = threading.RLock()


def _get_domain_dir(domain: str) -> Path:
    """Return the directory for a domain's FAISS index + metadata (not created here)."""
    collection_name = DOMAIN_COLLECTIONS.get(domain)
    if not collection_name:
        raise ValueError(f"Unknown domain: {domain}. Valid: {list(DOMAIN_COLLECTIONS.keys())}")
    return Path(FAISS_INDEX_DIR) / collection_name


def load_index(domain: str) -> tuple[faiss.Index, list[dict]]:
    """Load a domain's FAISS index + documents from disk, or create empty
    ones if nothing has been ingested yet. Returns (index, documents)."""
    domain_dir = _get_domain_dir(domain)
    index_path = domain_dir / "index.faiss"
    docs_path = domain_dir / "documents.json"

    with _index_lock:
        if index_path.exists() and docs_path.exists():
            index = faiss.read_index(str(index_path))
            documents = json.loads(docs_path.read_text(encoding="utf-8"))
        else:
            # IndexFlatL2 = brute-force search using Euclidean distance.
            # Simple and exact — no training needed. Good for up to ~100K vectors.
            index = faiss.IndexFlatL2(EMBEDDING_DIM)
            documents = []

    return index, documents


def _save_index(domain: str, index: faiss.Index, documents: list[dict]):
    """Persist the FAISS index and documents list to disk."""
    domain_dir = _get_domain_dir(domain)
    with _index_lock:
        domain_dir.mkdir(parents=True, exist_ok=True)
        faiss.write_index(index, str(domain_dir / "index.faiss"))
        (domain_dir / "documents.json").write_text(
            json.dumps(documents, ensure_ascii=False), encoding="utf-8"
        )


def _drop_filename(index: faiss.Index, documents: list[dict], filename: str):
    """Remove all chunks belonging to `filename` from the index.

    FAISS has no upsert, so re-uploading a file would otherwise append a
    duplicate copy of every chunk. We rebuild the index from the vectors we
    want to keep — `reconstruct_n` reads them back out of the flat index,
    so nothing has to be re-encoded.
    """
    keep = [i for i, doc in enumerate(documents)
            if doc["metadata"].get("filename") != filename]

    if len(keep) == len(documents):
        return index, documents  # nothing to remove

    new_index = faiss.IndexFlatL2(EMBEDDING_DIM)
    if keep:
        all_vectors = index.reconstruct_n(0, index.ntotal)
        new_index.add(np.array([all_vectors[i] for i in keep], dtype=np.float32))

    return new_index, [documents[i] for i in keep]


def embed_and_store(chunks: list[dict], domain: str) -> int:
    """Embed chunks and add them to the domain's FAISS index.
    Re-ingesting the same filename replaces its previous chunks.
    Returns the number of chunks stored."""
    if not chunks:
        return 0

    texts = [c["text"] for c in chunks]
    metadatas = [c["metadata"] for c in chunks]

    for m in metadatas:
        m["domain"] = domain

    # Encode all chunk texts into vectors (numpy array of shape [N, 384]).
    # This is the slow part, so it happens before taking the lock.
    embeddings = embedding_model.encode(texts, normalize_embeddings=True)
    embeddings = np.array(embeddings, dtype=np.float32)

    with _index_lock:
        index, documents = load_index(domain)
        index, documents = _drop_filename(index, documents, metadatas[0]["filename"])

        index.add(embeddings)

        # Store text + metadata in the parallel list (same order as vectors)
        for text, meta in zip(texts, metadatas):
            documents.append({"text": text, "metadata": meta})

        _save_index(domain, index, documents)

    return len(chunks)


def ingest_document(file_path: str, domain: str | None = None) -> dict:
    """Full pipeline: parse → detect domain (if not given) → chunk → embed → store.

    Domain detection happens here, after parsing, because it needs the
    document's actual text — the raw uploaded bytes of a PDF or DOCX are
    binary and carry no readable keywords.
    """
    filename = Path(file_path).name
    pages = parse_document(file_path)

    if not pages:
        message = "No readable text was found in this document."
        if Path(file_path).suffix.lower() == ".pdf":
            message += " Scanned PDFs (pages saved as images) aren't supported yet."
        raise DocumentError(message)

    if domain:
        detected_by = "user"
    else:
        sample_text = "\n".join(p["text"] for p in pages)[:5000]
        domain, detected_by = detect_domain(sample_text, filename)

    chunks = chunk_document(pages, filename)
    stored = embed_and_store(chunks, domain)
    return {
        "filename": filename,
        "domain": domain,
        "detected_by": detected_by,  # "user", "llm", or "keywords"
        "pages_parsed": len(pages),
        "chunks_stored": stored,
    }


# ── Listing and deleting ──


def list_documents() -> list[dict]:
    """Every stored document with its domain, chunk count and page count.
    Reads only the documents.json metadata files, not the vector indexes."""
    found = []
    with _index_lock:
        for domain in DOMAIN_COLLECTIONS:
            docs_path = _get_domain_dir(domain) / "documents.json"
            if not docs_path.exists():
                continue

            by_filename = {}
            for doc in json.loads(docs_path.read_text(encoding="utf-8")):
                meta = doc["metadata"]
                entry = by_filename.setdefault(
                    meta["filename"],
                    {"filename": meta["filename"], "domain": domain, "chunks": 0, "pages": 0},
                )
                entry["chunks"] += 1
                entry["pages"] = max(entry["pages"], meta.get("page", 1))

            found.extend(sorted(by_filename.values(), key=lambda d: d["filename"].lower()))
    return found


def delete_document(domain: str, filename: str) -> int:
    """Remove a document's chunks from its domain's index.
    Returns how many chunks were removed (0 means no such document)."""
    with _index_lock:
        index, documents = load_index(domain)
        new_index, kept = _drop_filename(index, documents, filename)
        removed = len(documents) - len(kept)
        if removed:
            _save_index(domain, new_index, kept)
    return removed
